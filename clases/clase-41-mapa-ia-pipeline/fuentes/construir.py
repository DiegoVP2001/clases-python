"""Builds the Clase 41 interactive HTML, exports the diagram PNGs and copies everything to the hub.

Run from anywhere:  py -3 construir.py
Source of truth: plantilla_interactivo.html + diagramas.py. Never edit the generated files by hand.
"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from diagramas import diagrama_mapa, diagrama_pipeline

AQUI = Path(__file__).resolve().parent
CLASE = AQUI.parent
HUB = CLASE.parents[1] / "sitio-hub"  # <repo>/sitio-hub
EDGE = Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe")
NOMBRE_HTML = "Clase 41 - Mapa IA - Interactivo.html"
NOMBRE_PPT = "Clase 41 - Mapa IA - Presentación.pptx"


def exportar_png(svg: str, destino: Path) -> None:
    """Renders a 1920x1080 SVG to PNG with headless Edge."""
    envoltorio = AQUI / "_render.html"
    envoltorio.write_text(
        "<!doctype html><meta charset='utf-8'><style>html,body{margin:0}svg{display:block}</style>" + svg,
        encoding="utf-8")
    subprocess.run([str(EDGE), "--headless=new", "--disable-gpu", "--hide-scrollbars",
                    f"--screenshot={destino}", "--window-size=1920,1080", envoltorio.as_uri()],
                   check=True, timeout=90)
    envoltorio.unlink(missing_ok=True)


def construir_ppt() -> Path:
    """One full-bleed PNG per slide (rendered from SVG) plus speaker notes, saved as .pptx."""
    from pptx import Presentation
    from pptx.util import Emu

    from diapositivas import diapositivas

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    # Temp PNGs go to the system temp folder, not the project: Edge leaves them with ACLs that
    # block deleting the folder, which broke `git add` when it lived under fuentes/.
    tmp = Path(tempfile.mkdtemp(prefix="clase41_slides_"))
    try:
        for n, (titulo_, svg, notas) in enumerate(diapositivas(), start=1):
            png = tmp / f"s{n:02d}.png"
            exportar_png(svg, png)
            slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
            pic = slide.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
            pic._element.nvPicPr.cNvPr.set("descr", titulo_)  # alt text for accessibility
            slide.notes_slide.notes_text_frame.text = notas
        destino = CLASE / NOMBRE_PPT
        prs.save(destino)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return destino


def main() -> int:
    svg1, svg2 = diagrama_mapa(), diagrama_pipeline()
    plantilla = (AQUI / "plantilla_interactivo.html").read_text(encoding="utf-8")
    html = plantilla.replace("{{SVG1}}", svg1).replace("{{SVG2}}", svg2)
    (CLASE / NOMBRE_HTML).write_text(html, encoding="utf-8")

    carpeta_png = CLASE / "diagramas"
    carpeta_png.mkdir(exist_ok=True)
    png1 = carpeta_png / "Diagrama 1 - El mapa de la IA.png"
    png2 = carpeta_png / "Diagrama 2 - El pipeline.png"
    exportar_png(svg1, png1)
    exportar_png(svg2, png2)

    destino = HUB / "clase-41"
    (destino / "diagramas").mkdir(parents=True, exist_ok=True)
    (destino / "index.html").write_text(html, encoding="utf-8")
    shutil.copy2(png1, destino / "diagramas" / "diagrama-1-mapa.png")
    shutil.copy2(png2, destino / "diagramas" / "diagrama-2-pipeline.png")
    print("HTML:", CLASE / NOMBRE_HTML)
    print("PNG :", png1.name, png1.stat().st_size, "bytes |", png2.name, png2.stat().st_size, "bytes")
    print("Hub :", destino)
    ppt = construir_ppt()
    print("PPT :", ppt.name, ppt.stat().st_size // 1024, "KB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
