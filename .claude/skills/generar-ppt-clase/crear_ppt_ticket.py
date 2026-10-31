#!/usr/bin/env python3
"""
crear_ppt_ticket.py — Generador del PPT *aparte* del Ticket de Salida.

Produce `Clase NN - Tema - Ticket de Salida.pptx`: las preguntas de
alternativas que Diego proyecta en la TV el día de la clase, con el mismo
diseño de marca que `Presentación.pptx` pero como archivo independiente.

Por qué es un archivo aparte (CLAUDE.md regla 17): `Presentación.pptx` se
sube a Classroom antes de dictar la clase, así que cualquier pregunta que
viva ahí queda expuesta con anticipación.

Estructura de slides:

    1.          Portada + reglas de la dinámica
    2..N+1      Una slide por pregunta (enunciado + código + 4 alternativas)
    N+2         Pantalla del Google Form (nadie responde antes de esta slide)
    N+3..       Revisión: una slide por pregunta con la correcta + justificación

Las slides de revisión van *después* de la del Form a propósito: el orden de
la dinámica evita el efecto arrastre (cada estudiante se compromete con su
respuesta antes de ver la de los demás).

Uso:
    python crear_ppt_ticket.py <ruta_spec.md> <ruta_salida.pptx>
    python crear_ppt_ticket.py <ruta_spec.md> <ruta_salida.pptx> --sin-revision

Requiere: pip install python-pptx
"""

import sys
import re
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Pt, Inches
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
except ImportError:
    print("ERROR: python-pptx no está instalado. Instálalo con:")
    print("    pip install python-pptx")
    sys.exit(1)

CARPETA_SKILL = Path(__file__).resolve().parent
sys.path.insert(0, str(CARPETA_SKILL))

from construir_plantilla import (
    aplicar_fondo_oscuro as _fondo_oscuro,
    agregar_franja_lateral as _franja_lateral,
    agregar_rectangulo as _rect,
    agregar_texto as _texto,
    C as _C, T as _T, E as _E,
    ANCHO_SLIDE as _W, ALTO_SLIDE as _H, MARGEN_X as _MX,
)

# El Form de registro recurrente del Ticket de Salida (CLAUDE.md regla 17).
URL_FORM = "https://forms.gle/sjRpbgmQzrpkEBsH9"

_ANCHO_UTIL = _W - 2 * _MX
_LETRAS = ("A", "B", "C", "D")


# =====================================================================
# PARSER DEL SPEC
# =====================================================================

def parsear_ticket(ruta_spec: Path) -> dict:
    """Extrae del spec el número de clase, el tema breve del Form y las
    preguntas del Ticket de Salida."""
    contenido = ruta_spec.read_text(encoding="utf-8")

    datos = {"numero": "?", "tema": "", "tema_breve_form": "", "preguntas": []}

    m = re.search(r"^#\s+Clase\s+([\w.]+)\s*[—\-–]\s*(.+)$", contenido, re.MULTILINE)
    if m:
        datos["numero"] = m.group(1).strip()
        datos["tema"] = m.group(2).strip()

    m = re.search(r"^\s*-\s*\*\*Tema breve \(Form\):\*\*\s*(.+)$", contenido, re.MULTILINE)
    if m:
        datos["tema_breve_form"] = m.group(1).strip()

    # Sección del Ticket: hasta el siguiente encabezado de nivel 2 o 3.
    m = re.search(r"^###\s+\d+\.\s+Ticket de Salida.*$", contenido, re.MULTILINE)
    if not m:
        raise SystemExit("ERROR: el spec no tiene una sección '### N. Ticket de Salida'.")
    resto = contenido[m.end():]
    corte = re.search(r"^#{2,3}\s", resto, re.MULTILINE)
    seccion = resto[:corte.start()] if corte else resto

    # Cada pregunta arranca en **Pregunta N:**
    trozos = re.split(r"^\*\*Pregunta\s+\d+:?\*\*\s*$", seccion, flags=re.MULTILINE)
    for trozo in trozos[1:]:
        datos["preguntas"].append(_parsear_pregunta(trozo))

    if not datos["preguntas"]:
        raise SystemExit("ERROR: no se encontró ninguna '**Pregunta N:**' en el Ticket.")
    return datos


def _parsear_pregunta(texto: str) -> dict:
    preg = {"enunciado": "", "codigo": "", "alternativas": [],
            "correcta": "", "justificacion": ""}

    m = re.search(r"```(?:python)?\n(.*?)```", texto, re.DOTALL)
    if m:
        preg["codigo"] = m.group(1).rstrip("\n")
        texto = texto[:m.start()] + "\n" + texto[m.end():]

    for linea in texto.split("\n"):
        linea = linea.strip()
        if not linea:
            continue
        m = re.match(r"^-\s*([ABCD])[:.]\s*(.+)$", linea)
        if m:
            preg["alternativas"].append((m.group(1), m.group(2).strip()))
            continue
        m = re.match(r"^\*\*Respuesta correcta:?\*\*\s*(.+)$", linea)
        if m:
            preg["correcta"] = m.group(1).strip().strip("`")
            continue
        m = re.match(r"^\*\*Justificación:?\*\*\s*(.+)$", linea)
        if m:
            preg["justificacion"] = m.group(1).strip()
            continue
        if linea.startswith(">") or linea.startswith("**"):
            continue  # notas de conducción del spec: no se proyectan
        preg["enunciado"] += (" " if preg["enunciado"] else "") + linea

    return preg


# =====================================================================
# TEXTO RICO (inline code + negrita)
# =====================================================================

_RE_TOKEN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")


def _limpiar(s: str) -> str:
    """Quita escapes que solo existen para el markdown de Jupyter."""
    return (s or "").replace("\\$", "$").replace("\\*", "*")


def _texto_rico(slide, x, y, w, h, parrafos, tamano, color="blanco",
                fuente="Calibri", alineacion=PP_ALIGN.LEFT,
                anchor=MSO_ANCHOR.TOP, interlineado=1.0):
    """Caja de texto con `código inline` en Consolas verdoso y **negrita**.

    `parrafos` es una lista de strings; cada uno es un párrafo.
    """
    cuadro = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = cuadro.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)

    for i, parrafo in enumerate(parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alineacion
        p.line_spacing = interlineado
        if i > 0:
            p.space_before = Pt(tamano * 0.35)
        for token in _RE_TOKEN.split(_limpiar(parrafo)):
            if not token:
                continue
            run = p.add_run()
            if token.startswith("`") and token.endswith("`"):
                run.text = token[1:-1]
                run.font.name = "Consolas"
                run.font.color.rgb = _C["codigo_color"]
                run.font.size = Pt(tamano - 1)
            elif token.startswith("**") and token.endswith("**"):
                run.text = token[2:-2]
                run.font.name = fuente
                run.font.bold = True
                run.font.color.rgb = _C[color] if isinstance(color, str) else color
                run.font.size = Pt(tamano)
            else:
                run.text = token
                run.font.name = fuente
                run.font.color.rgb = _C[color] if isinstance(color, str) else color
                run.font.size = Pt(tamano)
    return cuadro


def _texto_plano(s: str) -> str:
    """Versión sin marcas, para medir cuánto ocupa."""
    return re.sub(r"[`*]", "", _limpiar(s))


def _estimar_lineas(texto: str, tamano_pt: int, ancho_in: float,
                    ratio: float = 0.50) -> int:
    if not texto:
        return 0
    ancho_util_pt = max((ancho_in - 0.20) * 72.0, 1.0)
    chars_por_linea = max(int(ancho_util_pt / (ratio * tamano_pt)), 10)
    total = 0
    for linea in texto.split("\n"):
        total += max(1, -(-len(linea) // chars_por_linea))
    return total


# =====================================================================
# CHROME DEL SLIDE
# =====================================================================

def _slide_base(prs, numero_clase, seccion):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fondo_oscuro(slide)
    _franja_lateral(slide)
    _rect(slide, 0, 0, _W, 0.55, _C["tabla_header"])
    _texto(slide, _MX - 0.15, 0.13, 6, 0.35,
           f"🐍 Python · Clase {numero_clase}", tamano=_T["barra_navegacion"],
           color="gris_secundario", fuente="Consolas",
           anchor_vertical=MSO_ANCHOR.MIDDLE)
    _texto(slide, _W - 4.15, 0.13, 3.5, 0.35, seccion,
           tamano=_T["barra_navegacion"], color="ambar", fuente="Consolas",
           negrita=True, alineacion=PP_ALIGN.RIGHT,
           anchor_vertical=MSO_ANCHOR.MIDDLE)
    return slide


def _bloque_codigo(slide, x, y, w, h, codigo, tam):
    _rect(slide, x, y, w, h, _C["terminal_bg"])
    _rect(slide, x, y, w, 0.10, _C["turquesa"])
    _texto(slide, x + 0.15, y + 0.16, w - 0.30, h - 0.24, codigo,
           tamano=tam, color="codigo_color", fuente="Consolas")


def _badge(slide, x, y, d, letra, color_relleno="ambar"):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                                   Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _C[color_relleno]
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = letra
    run.font.name = "Calibri"
    run.font.bold = True
    run.font.size = Pt(int(d * 44))
    run.font.color.rgb = _C["fondo_oscuro"]


# =====================================================================
# SLIDES
# =====================================================================

def slide_portada(prs, datos):
    slide = _slide_base(prs, datos["numero"], "TICKET DE SALIDA")

    _texto(slide, _MX, 1.05, _ANCHO_UTIL, 0.95,
           f"{_E['ticket']}  Ticket de Salida", tamano=44, color="ambar",
           negrita=True)
    _texto(slide, _MX, 2.00, _ANCHO_UTIL, 0.50,
           f"Clase {datos['numero']} — {datos['tema']}", tamano=22,
           color="blanco")

    n = len(datos["preguntas"])
    alto = 3.35
    _rect(slide, _MX, 2.80, _ANCHO_UTIL, alto, _C["fondo_bloque"])
    _rect(slide, _MX, 2.80, 0.08, alto, _C["turquesa"])
    _texto(slide, _MX + 0.28, 2.94, _ANCHO_UTIL - 0.5, 0.42,
           f"{_E['idea']}  Cómo lo vamos a hacer", tamano=_T["titulo_bloque"],
           color="turquesa", negrita=True)
    _texto_rico(slide, _MX + 0.28, 3.45, _ANCHO_UTIL - 0.6, 2.55, [
        f"1.  Son {n} preguntas. Las vemos **una por una**.",
        "2.  Decide tu respuesta **en silencio**: no la digas ni la comentes.",
        "3.  Anota la letra que elegiste (**A**, **B**, **C** o **D**).",
        f"4.  Recién al terminar la pregunta {n} completamos el Formulario.",
        "5.  Después de enviarlo, revisamos juntos las respuestas.",
    ], tamano=22, interlineado=1.05)

    _texto(slide, _MX, 6.50, _ANCHO_UTIL, 0.45,
           "Nadie responde el Formulario antes de tiempo: primero te "
           "comprometes con tu respuesta.", tamano=17, color="gris_claro")
    return slide


# — Solver de alto para las slides de pregunta —
_Y_INI = 1.28
_Y_FIN = 7.30
_GAP = 0.20
_GAP_FILAS = 0.14
_ESCALA_COD = (20, 18, 16, 14)
_ESCALA_ALT = (24, 22, 20, 18, 16)
_ANCHO_BADGE = 0.95   # badge + separación, dentro de la celda


def _ancho_celda(cols):
    return (_ANCHO_UTIL - 0.30) / 2 if cols == 2 else _ANCHO_UTIL


def _medir(preg, tam_cod, tam_alt, tam_enun, cols):
    lin_e = _estimar_lineas(_texto_plano(preg["enunciado"]), tam_enun,
                            _ANCHO_UTIL - 0.5)
    h_enun = lin_e * (tam_enun * 1.25 / 72.0) + 0.22

    h_cod = 0.0
    if preg["codigo"]:
        n = len(preg["codigo"].split("\n"))
        h_cod = n * (tam_cod * 1.20 / 72.0) + 0.32

    ancho_txt = _ancho_celda(cols) - _ANCHO_BADGE
    alturas_fila = []
    alts = preg["alternativas"]
    for i in range(0, len(alts), cols):
        grupo = alts[i:i + cols]
        lineas = max(_estimar_lineas(_texto_plano(t), tam_alt, ancho_txt)
                     for _, t in grupo)
        alturas_fila.append(lineas * (tam_alt * 1.22 / 72.0) + 0.28)

    total = h_enun + h_cod + sum(alturas_fila)
    total += _GAP * (1 + (1 if h_cod else 0))
    total += _GAP_FILAS * (len(alturas_fila) - 1)
    return total, h_enun, h_cod, alturas_fila


# Peso de cada punto de fuente sacrificado. El código pesa más: es lo que hay
# que leer desde el fondo de la sala, y es el objeto de la pregunta. Las
# alternativas son frases cortas, aguantan un par de puntos menos sin
# volverse ilegibles (calibrado con la Clase 19.5, donde Diego bajó a mano
# las alternativas a ~20pt para darle aire al terminal).
_PESO_COD = 2.0
_PESO_ALT = 1.0


def _mejor_para_layout(preg, cols, area, escala_cod):
    """Mejor combinación (código, alternativas) para un layout dado.

    Se elige por costo, y el código pesa el doble: cuando hay que sacrificar
    tamaño, se sacrifica primero el de las alternativas. Devuelve None si
    nada cabe.
    """
    mejor = None
    for tam_cod in escala_cod:
        for tam_alt in _ESCALA_ALT:
            tam_enun = max(20, min(24, tam_alt + 2))
            medida = _medir(preg, tam_cod, tam_alt, tam_enun, cols)
            if medida[0] > area:
                continue
            costo = (20 - tam_cod) * _PESO_COD + (24 - tam_alt) * _PESO_ALT
            if mejor is None or costo < mejor[0]:
                mejor = (costo, tam_cod, tam_alt, tam_enun) + medida[1:]
    return mejor


def _resolver_pregunta(preg):
    """Elige layout y tamaños de fuente.

    Se prueban los dos layouts posibles de alternativas —grilla 2x2 y
    columna única— y gana el de menor costo: con textos cortos rinde más la
    grilla, y con textos largos (una alternativa de tres líneas) rinde más la
    columna única, que aprovecha todo el ancho.
    """
    area = _Y_FIN - _Y_INI
    escala_cod = tuple(t for t in _ESCALA_COD
                       if _cabe_ancho(preg["codigo"], t, _ANCHO_UTIL)) \
        or (_ESCALA_COD[-1],)

    candidatos = []
    for cols in (2, 1):
        mejor = _mejor_para_layout(preg, cols, area, escala_cod)
        if mejor:
            candidatos.append((mejor, cols))

    if candidatos:
        # Menor costo gana; a igualdad, el que deje el código más grande.
        (mejor, cols) = min(candidatos, key=lambda c: (c[0][0], -c[0][1]))
        _, tam_cod, tam_alt, tam_enun, h_enun, h_cod, filas = mejor
        return tam_cod, tam_alt, tam_enun, h_enun, h_cod, filas, cols, True

    # Nada cabe: se usa el piso de la escala y se avisa.
    tam_cod, tam_alt, tam_enun = escala_cod[-1], _ESCALA_ALT[-1], 20
    medida = _medir(preg, tam_cod, tam_alt, tam_enun, 1)
    return (tam_cod, tam_alt, tam_enun) + medida[1:] + (1, False)


def _cabe_ancho(codigo, tam, ancho_in):
    if not codigo:
        return True
    max_chars = max(len(l) for l in codigo.split("\n"))
    return max_chars * 0.60 * tam <= (ancho_in - 0.30) * 72.0


def slide_pregunta(prs, datos, idx):
    preg = datos["preguntas"][idx]
    total = len(datos["preguntas"])
    slide = _slide_base(prs, datos["numero"], "TICKET DE SALIDA")

    _texto(slide, _MX, 0.62, _ANCHO_UTIL, 0.60,
           f"{_E['ticket']}  Pregunta {idx + 1} de {total}", tamano=30,
           color="ambar", negrita=True)

    (tam_cod, tam_alt, tam_enun, h_enun, h_cod,
     alturas_fila, cols, ok) = _resolver_pregunta(preg)
    if not ok:
        print(f"   ⚠️  Pregunta {idx + 1}: no cabe ni con la fuente más chica "
              f"— acorta el texto en el spec")

    y = _Y_INI

    _rect(slide, _MX, y, _ANCHO_UTIL, h_enun, _C["fondo_bloque"])
    _rect(slide, _MX, y, 0.08, h_enun, _C["turquesa"])
    _texto_rico(slide, _MX + 0.28, y + 0.12, _ANCHO_UTIL - 0.5, h_enun - 0.20,
                [preg["enunciado"]], tamano=tam_enun,
                anchor=MSO_ANCHOR.MIDDLE)
    y += h_enun + _GAP

    if h_cod:
        _bloque_codigo(slide, _MX, y, _ANCHO_UTIL, h_cod, preg["codigo"], tam_cod)
        y += h_cod + _GAP

    ancho_celda = _ancho_celda(cols)
    ancho_txt = ancho_celda - _ANCHO_BADGE
    alts = preg["alternativas"]
    for fila, i in enumerate(range(0, len(alts), cols)):
        h = alturas_fila[fila]
        for col, (letra, txt) in enumerate(alts[i:i + cols]):
            x = _MX + col * (ancho_celda + 0.30)
            _rect(slide, x, y, ancho_celda, h, _C["fondo_bloque"])
            d = min(0.52, h - 0.16)
            _badge(slide, x + 0.18, y + (h - d) / 2, d, letra)
            _texto_rico(slide, x + 0.18 + d + 0.16, y + 0.08,
                        ancho_txt, h - 0.16, [txt], tamano=tam_alt,
                        anchor=MSO_ANCHOR.MIDDLE)
        y += h + _GAP_FILAS
    return slide


def slide_formulario(prs, datos):
    slide = _slide_base(prs, datos["numero"], "TICKET DE SALIDA")

    _texto(slide, _MX, 0.72, _ANCHO_UTIL, 0.70,
           "📝  Ahora sí: completa el Formulario", tamano=34, color="menta",
           negrita=True)

    _rect(slide, _MX, 1.65, _ANCHO_UTIL, 1.20, _C["terminal_bg"])
    _rect(slide, _MX, 1.65, _ANCHO_UTIL, 0.10, _C["menta"])
    _texto(slide, _MX, 1.95, _ANCHO_UTIL, 0.75, URL_FORM, tamano=32,
           color="codigo_color", fuente="Consolas", alineacion=PP_ALIGN.CENTER)

    alto = 3.55
    _rect(slide, _MX, 3.10, _ANCHO_UTIL, alto, _C["fondo_bloque"])
    _rect(slide, _MX, 3.10, 0.08, alto, _C["menta"])
    _texto(slide, _MX + 0.28, 3.24, _ANCHO_UTIL - 0.5, 0.42,
           f"{_E['idea']}  Qué escribir", tamano=_T["titulo_bloque"],
           color="menta", negrita=True)

    tema = datos["tema_breve_form"] or datos["tema"]
    _texto_rico(slide, _MX + 0.28, 3.78, _ANCHO_UTIL - 0.6, 2.70, [
        "**Nombre:** tu nombre completo.",
        f"**Tema de la clase de hoy:** `{tema}`",
        f"**Respuestas:** marca la letra que anotaste en cada una de las "
        f"{len(datos['preguntas'])} preguntas.",
        "Envía. Cuando estén todos, revisamos las respuestas juntos.",
    ], tamano=22, interlineado=1.05)
    return slide


def slide_revision(prs, datos, idx):
    preg = datos["preguntas"][idx]
    slide = _slide_base(prs, datos["numero"], "REVISIÓN")

    _texto(slide, _MX, 0.62, _ANCHO_UTIL, 0.60,
           f"✅  Pregunta {idx + 1} — Respuesta", tamano=30, color="menta",
           negrita=True)

    _texto_rico(slide, _MX, 1.30, _ANCHO_UTIL, 0.60, [preg["enunciado"]],
                tamano=19, color="gris_claro")

    correcta = preg["correcta"]
    texto_alt = next((t for l, t in preg["alternativas"] if l == correcta), "")

    h_alt = max(0.95, _estimar_lineas(_texto_plano(texto_alt), 24,
                                      _ANCHO_UTIL - 1.3) * 0.42 + 0.50)
    y = 2.05
    _rect(slide, _MX, y, _ANCHO_UTIL, h_alt, _C["verde_fondo"], _C["verde"])
    d = 0.60
    _badge(slide, _MX + 0.22, y + (h_alt - d) / 2, d, correcta,
           color_relleno="verde")
    _texto_rico(slide, _MX + 1.05, y + 0.10, _ANCHO_UTIL - 1.35, h_alt - 0.20,
                [texto_alt], tamano=24, anchor=MSO_ANCHOR.MIDDLE)

    y += h_alt + 0.30
    alto = _Y_FIN - y
    _rect(slide, _MX, y, _ANCHO_UTIL, alto, _C["fondo_bloque"])
    _rect(slide, _MX, y, 0.08, alto, _C["ambar"])
    _texto(slide, _MX + 0.28, y + 0.14, _ANCHO_UTIL - 0.5, 0.42,
           f"{_E['idea']}  Por qué", tamano=_T["titulo_bloque"], color="ambar",
           negrita=True)
    _texto_rico(slide, _MX + 0.28, y + 0.68, _ANCHO_UTIL - 0.6, alto - 0.82,
                [preg["justificacion"]], tamano=22, interlineado=1.05)
    return slide


# =====================================================================
# MAIN
# =====================================================================

def generar(ruta_spec: Path, ruta_salida: Path, con_revision=True):
    datos = parsear_ticket(ruta_spec)

    prs = Presentation()
    prs.slide_width = Inches(_W)
    prs.slide_height = Inches(_H)

    slide_portada(prs, datos)
    for i in range(len(datos["preguntas"])):
        slide_pregunta(prs, datos, i)
    slide_formulario(prs, datos)
    if con_revision:
        for i in range(len(datos["preguntas"])):
            slide_revision(prs, datos, i)

    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(ruta_salida))

    print(f"✅ {ruta_salida.name}")
    print(f"   {len(prs.slides)} slides · {len(datos['preguntas'])} preguntas")
    letras = ", ".join(f"P{i+1}={p['correcta']}"
                       for i, p in enumerate(datos["preguntas"]))
    print(f"   Respuestas correctas: {letras}")


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) < 2:
        print(__doc__)
        sys.exit(1)
    generar(Path(args[0]), Path(args[1]),
            con_revision="--sin-revision" not in sys.argv)


if __name__ == "__main__":
    main()
